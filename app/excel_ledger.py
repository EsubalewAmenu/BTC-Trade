import shutil
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


TRADE_HEADERS = [
    "Trade ID", "Status", "Side", "Opened UTC", "Closed UTC", "Entry Price",
    "Exit Price", "Quantity BTC", "Notional USDT", "Stop Price", "Target Price",
    "Entry Fee", "Exit Fee", "Gross PnL", "Net PnL", "Exit Reason", "Hold Minutes",
    "LLM Confidence", "LLM Rationale", "Signal Candle UTC", "Balance After",
]
EVENT_HEADERS = [
    "Timestamp UTC", "Type", "Candle UTC", "Price", "Rule Signal", "LLM Action",
    "Confidence", "Result", "Account Available USDT", "Reason",
]


class ExcelLedger:
    def __init__(self, data_dir: Path, start_balance: float, settings_rows: list):
        self.path = Path(data_dir) / "paper_trading.xlsx"
        self.path.parent.mkdir(parents=True, exist_ok=True)
        if not self.path.exists():
            template = Path(__file__).with_name("paper_trading_template.xlsx")
            if template.exists():
                shutil.copy2(template, self.path)
                self._sync_template(start_balance, settings_rows)
            else:
                self._create(start_balance, settings_rows)

    def _sync_template(self, start_balance, settings_rows):
        wb = load_workbook(self.path)
        wb["Dashboard"]["B5"] = start_balance
        sheet = wb["Settings"]
        for row_number, row in enumerate(settings_rows, start=4):
            sheet.cell(row_number, 1).value = row[0]
            sheet.cell(row_number, 2).value = row[1]
        self._save(wb)

    def _create(self, start_balance, settings_rows):
        wb = Workbook()
        dashboard = wb.active
        dashboard.title = "Dashboard"
        trades = wb.create_sheet("Trades")
        events = wb.create_sheet("Events")
        settings = wb.create_sheet("Settings")
        dashboard.append(["BTCUSDT Intraday Paper Trading"])
        dashboard.append(["Metric", "Value"])
        dashboard.append(["Starting Balance", start_balance])
        dashboard.append(["Current Balance", "=IFERROR(LOOKUP(2,1/(Trades!U2:U10001<>\"\"),Trades!U2:U10001),B3)"])
        dashboard.append(["Closed Trades", '=COUNTIF(Trades!B2:B10001,"CLOSED")'])
        dashboard.append(["Wins", '=COUNTIF(Trades!P2:P10001,"TARGET")'])
        dashboard.append(["Win Rate", '=IFERROR(B6/B5,0)'])
        dashboard.append(["Net PnL", '=SUM(Trades!O2:O10001)'])
        dashboard.append(["Profit Factor", '=IFERROR(SUMIF(Trades!O2:O10001,">0",Trades!O2:O10001)/ABS(SUMIF(Trades!O2:O10001,"<0",Trades!O2:O10001)),0)'])
        trades.append(TRADE_HEADERS)
        events.append(EVENT_HEADERS)
        settings.append(["Setting", "Value"])
        for row in settings_rows:
            settings.append(row)
        for sheet in wb.worksheets:
            sheet.freeze_panes = "A2"
            sheet.sheet_view.showGridLines = False
            sheet["A1"].font = Font(bold=True, color="FFFFFF")
            for cell in sheet[1]:
                cell.fill = PatternFill("solid", fgColor="17365D")
                cell.font = Font(bold=True, color="FFFFFF")
                cell.alignment = Alignment(vertical="center")
            for column in sheet.columns:
                letter = column[0].column_letter
                width = min(max(len(str(cell.value or "")) for cell in column) + 2, 42)
                sheet.column_dimensions[letter].width = max(width, 12)
        dashboard["B3"].number_format = '$#,##0.00;[Red]($#,##0.00)'
        dashboard["B4"].number_format = '$#,##0.00;[Red]($#,##0.00)'
        dashboard["B7"].number_format = "0.0%"
        dashboard["B8"].number_format = '$#,##0.00;[Red]($#,##0.00)'
        self._save(wb)

    def _save(self, wb):
        temporary = self.path.with_name(f"{self.path.stem}.tmp.xlsx")
        wb.save(temporary)
        temporary.replace(self.path)

    def record_event(self, event_type, analysis, decision, account, result):
        wb = load_workbook(self.path)
        sheet = wb["Events"]
        from datetime import datetime, timezone
        sheet.append([
            datetime.now(timezone.utc).isoformat(), event_type, analysis.candle_time,
            analysis.close, analysis.rule_signal, decision.action, decision.confidence,
            result, account.get("available_balance", ""), decision.rationale,
        ])
        self._save(wb)

    def open_trade(self, position, analysis):
        wb = load_workbook(self.path)
        sheet = wb["Trades"]
        sheet.append([
            position.trade_id, "OPEN", position.side, position.opened_at, "",
            position.entry_price, "", position.quantity, position.notional,
            position.stop_price, position.target_price, position.entry_fee, "", "", "",
            "", "", position.confidence, position.rationale, analysis.candle_time, "",
        ])
        self._save(wb)

    def close_trade(self, trade_id, result):
        wb = load_workbook(self.path)
        sheet = wb["Trades"]
        for row in range(2, sheet.max_row + 1):
            if sheet.cell(row, 1).value == trade_id:
                updates = {
                    2: "CLOSED", 5: result["closed_at"], 7: result["exit_price"],
                    13: result["exit_fee"], 14: result["gross_pnl"], 15: result["net_pnl"],
                    16: result["exit_reason"], 17: result["hold_minutes"], 21: result["balance"],
                }
                for column, value in updates.items():
                    sheet.cell(row, column).value = value
                break
        self._save(wb)
